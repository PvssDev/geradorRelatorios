# -*- coding: utf-8 -*-
"""
Serviço responsável pelo processamento e renderização do
Quadro 1 – Demonstrativo da Elegibilidade do Trecho Fiscalizado para Indicação de NC
utilizado exclusivamente nas fiscalizações da CRA.
"""

import io
import openpyxl
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def extrair_dados_elegibilidade(planilha_input) -> dict | None:
    """
    Lê a planilha de elegibilidade enviada pelo usuário (caminho, BytesIO ou UploadedFile)
    e extrai os cabeçalhos e registros organizados por pista/sentido.
    """
    if not planilha_input:
        return None

    # Normalizar input para stream em memória (LGPD / BytesIO)
    if hasattr(planilha_input, "read"):
        if hasattr(planilha_input, "seek"):
            planilha_input.seek(0)
        file_bytes = io.BytesIO(planilha_input.read())
        if hasattr(planilha_input, "seek"):
            planilha_input.seek(0)
        wb = openpyxl.load_workbook(file_bytes, data_only=True)
    elif isinstance(planilha_input, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(planilha_input), data_only=True)
    else:
        wb = openpyxl.load_workbook(planilha_input, data_only=True)

    # 1. Localizar a aba correta (preferência para 'TDE' ou termos de elegibilidade)
    target_sheet = None
    for name in ["TDE", "ELEGIBILIDADE", "DEMONSTRATIVO"]:
        for sname in wb.sheetnames:
            if name in sname.upper():
                target_sheet = wb[sname]
                break
        if target_sheet:
            break

    if not target_sheet:
        for sname in wb.sheetnames:
            ws = wb[sname]
            found = False
            for r in range(1, min(10, ws.max_row + 1)):
                for c in range(1, min(10, ws.max_column + 1)):
                    if "ELEG" in str(ws.cell(r, c).value or "").upper():
                        target_sheet = ws
                        found = True
                        break
                if found:
                    break
            if target_sheet:
                break

    if not target_sheet:
        target_sheet = wb.active

    # 2. Localizar a linha de cabeçalho
    header_row_idx = None
    for r in range(1, min(10, target_sheet.max_row + 1)):
        row_str = " ".join([str(target_sheet.cell(r, c).value or "") for c in range(1, 10)])
        if "TRECHO" in row_str.upper() and "ELEG" in row_str.upper():
            header_row_idx = r
            break

    if not header_row_idx:
        header_row_idx = 2

    # Extrair e normalizar os títulos das colunas
    c1_h = target_sheet.cell(header_row_idx, 1).value or "TRECHO ELEGÍVEL"
    c2_h = target_sheet.cell(header_row_idx, 2).value or "SUBTRECHO"
    c3_h = target_sheet.cell(header_row_idx, 3).value or "IGG"
    c4_h = str(target_sheet.cell(header_row_idx, 4).value or "").strip()
    c5_h = target_sheet.cell(header_row_idx, 5).value or "IRI"
    c6_h = str(target_sheet.cell(header_row_idx, 6).value or "").strip()

    def expand_relatorio(val_str: str) -> str:
        if not val_str:
            return ""
        s = val_str.strip()
        if s.startswith("(Rel ") and not s.startswith("(Relatório Anual "):
            return s.replace("(Rel ", "(Relatório Anual ")
        return s

    hdr_col0 = "TRECHO ELEGÍVEL\n(KMinicial-KMfinal)"
    hdr_col1 = str(c2_h).strip()
    c4_exp = expand_relatorio(c4_h)
    c6_exp = expand_relatorio(c6_h)
    hdr_col2 = f"{c3_h}\n{c4_exp}" if c4_exp else str(c3_h).strip()
    hdr_col3 = f"{c5_h}\n{c6_exp}" if c6_exp else str(c5_h).strip()

    headers = [hdr_col0, hdr_col1, hdr_col2, hdr_col3]

    # 3. Iterar pelas linhas extraindo seções (sentidos de pista) e linhas de dados
    sections = []
    current_track = None
    current_items = []

    for r in range(header_row_idx + 1, target_sheet.max_row + 1):
        c1 = target_sheet.cell(r, 1).value
        c2 = target_sheet.cell(r, 2).value
        c3 = target_sheet.cell(r, 3).value
        c4 = target_sheet.cell(r, 4).value
        c5 = target_sheet.cell(r, 5).value
        c6 = target_sheet.cell(r, 6).value

        if c1 is None and c2 is None and c3 is None:
            continue

        c1_str = str(c1 or "").strip()
        if not c1_str:
            continue

        # Se c2 estiver vazio, trata-se do cabeçalho da seção de pista/sentido
        if c2 is None or str(c2).strip() == "":
            if current_track or current_items:
                sections.append((current_track, current_items))
            current_track = c1_str
            current_items = []
        else:
            subtrecho = str(c2 or "").strip()

            # Formatação do IGG
            if isinstance(c3, float):
                igg_val = f"{c3:.2f}".replace(".", ",")
            else:
                igg_val = str(c3 or "").strip()
            c4_str = str(c4 or "").strip()
            igg_cell = f"{igg_val}\n{c4_str}" if c4_str else igg_val

            # Formatação do IRI
            if isinstance(c5, float):
                iri_val = f"{c5:.2f}".replace(".", ",")
            else:
                iri_val = str(c5 or "").strip()
            c6_str = str(c6 or "").strip()
            iri_cell = f"{iri_val}\n{c6_str}" if c6_str else iri_val

            current_items.append((c1_str, subtrecho, igg_cell, iri_cell))

    if current_track or current_items:
        sections.append((current_track, current_items))

    return {
        "headers": headers,
        "sections": sections
    }


def renderizar_quadro_elegibilidade(doc, dados_elegibilidade: dict) -> None:
    """
    Renderiza o Quadro 1 – Demonstrativo da Elegibilidade no documento Word
    com a estilização estrita idêntica ao modelo oficial (.docx da pasta ref).
    """
    if not dados_elegibilidade or not dados_elegibilidade.get("sections"):
        return

    headers = dados_elegibilidade.get("headers") or [
        "TRECHO ELEGÍVEL\n(KMinicial-KMfinal)",
        "SUBTRECHO",
        "IGG\n(Relatório Anual 01/2025)",
        "IRI\n(Relatório Anual 01/2025)"
    ]
    sections = dados_elegibilidade.get("sections", [])

    # Título do Quadro 1
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)

    r1 = p.add_run("Quadro 1")
    r1.bold = True
    r1.font.name = 'Aptos'
    r1.font.size = Pt(11)

    r2 = p.add_run(" – Demonstrativo da Elegibilidade do Trecho Fiscalizado para Indicação de NC")
    r2.font.name = 'Aptos'
    r2.font.size = Pt(11)

    # Cálculo do total de linhas: 1 cabeçalho + (1 linha de título + N linhas de dados por seção)
    total_rows = 1 + sum(1 + len(items) for _, items in sections)
    table = doc.add_table(rows=total_rows, cols=4)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Definir larguras de coluna idênticas ao modelo de referência
    col_widths = [Inches(1.80), Inches(1.10), Inches(1.90), Inches(1.90)]
    for i, col in enumerate(table.columns):
        col.width = col_widths[i]

    # 1. Cabeçalho Principal (Linha 0)
    hdr_row = table.rows[0]
    for c_idx, text in enumerate(headers):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths[c_idx]
        cell.text = text
        shd = parse_xml(r'<w:shd {} w:fill="DDDDDD"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shd)
        
        # Centralização vertical e horizontal
        valign = parse_xml(r'<w:vAlign {} w:val="center"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(valign)

        for p_cell in cell.paragraphs:
            p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cell.paragraph_format.space_before = Pt(2)
            p_cell.paragraph_format.space_after = Pt(2)
            for run in p_cell.runs:
                run.font.name = 'Aptos'
                run.font.size = Pt(10)
                run.bold = True

    # 2. Preenchimento de Seções de Pista e Registros
    curr_row = 1
    for track_title, items in sections:
        # Linha mesclada de Sentido da Pista
        row_track = table.rows[curr_row]
        a = row_track.cells[0]
        b = row_track.cells[3]
        merged_cell = a.merge(b)
        merged_cell.text = track_title or ""
        
        shd_track = parse_xml(r'<w:shd {} w:fill="FFFFCC"/>'.format(nsdecls('w')))
        merged_cell._tc.get_or_add_tcPr().append(shd_track)
        valign = parse_xml(r'<w:vAlign {} w:val="center"/>'.format(nsdecls('w')))
        merged_cell._tc.get_or_add_tcPr().append(valign)

        for p_cell in merged_cell.paragraphs:
            p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p_cell.paragraph_format.space_before = Pt(2)
            p_cell.paragraph_format.space_after = Pt(2)
            for run in p_cell.runs:
                run.font.name = 'Aptos'
                run.font.size = Pt(10)
                run.bold = False

        curr_row += 1

        # Linhas de Dados do Trecho
        for col0, col1, col2, col3 in items:
            row_data = table.rows[curr_row]
            row_data.cells[0].text = str(col0)
            row_data.cells[1].text = str(col1)
            row_data.cells[2].text = str(col2)
            row_data.cells[3].text = str(col3)

            for idx_c, cell in enumerate(row_data.cells):
                cell.width = col_widths[idx_c]
                valign = parse_xml(r'<w:vAlign {} w:val="center"/>'.format(nsdecls('w')))
                cell._tc.get_or_add_tcPr().append(valign)

                for p_cell in cell.paragraphs:
                    p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cell.paragraph_format.space_before = Pt(2)
                    p_cell.paragraph_format.space_after = Pt(2)
                    for run in p_cell.runs:
                        run.font.name = 'Aptos'
                        run.font.size = Pt(10)
                        run.bold = False

            curr_row += 1

    doc.add_paragraph()  # Espaçamento após a tabela
