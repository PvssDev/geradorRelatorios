# -*- coding: utf-8 -*-
import sys
import os
import io
import pandas as pd
from docx import Document

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src')))

from report import gerar_relatorio
from services.data_service import gerar_planilha_excel_buffer


def test_gerar_relatorio_crc_pipeline(tmp_path):
    temp_fiscalizacoes = [
        {
            "ID da Fiscalização": "2026-001",
            "Data": "27/05/2026",
            "Hora": "10:00",
            "Cidade": "Cabo de Santo Agostinho",
            "Local": "Sistema Viário do Paiva",
            "Pessoal Responsável": "Pedro Souza",
            "Coordenador": "Coordenador ARPE",
            "Contrato": "CRC-01",
            "Período": "2026",
            "Relatório Gerado": False
        }
    ]

    temp_nc = [
        {
            "ID da Fiscalização": "2026-001",
            "Nº": 1,
            "Não Conformidade": "FI",
            "Identificação": "CRC.SH015.0646+0648/2026.001",
            "Foto": "foto1.jpg",
            "Direção (faixa)": "Faixa 1",
            "Fundamento da infração": "PER Anexo IV",
            "Determinação": "Reparar fissuras"
        },
        {
            "ID da Fiscalização": "2026-001",
            "Nº": 2,
            "Não Conformidade": "Tachões soltos",
            "Identificação": "CRC.SH015.0646+0648/2026.002",
            "Foto": "foto2.jpg",
            "Direção (faixa)": "Faixa 2",
            "Fundamento da infração": "PER Anexo IV",
            "Determinação": "Substituir tachões"
        }
    ]

    excel_buf = gerar_planilha_excel_buffer(temp_fiscalizacoes, temp_nc)

    reports_dir = str(tmp_path / "reports")
    fotos_dir = str(tmp_path / "fotos")
    os.makedirs(reports_dir, exist_ok=True)
    os.makedirs(fotos_dir, exist_ok=True)

    arquivos, _ = gerar_relatorio(
        caminho_planilha=excel_buf,
        fotos_dir=fotos_dir,
        relatorios_dir=reports_dir,
        gerar_todos=True,
        tipo_relatorio="CRC"
    )

    assert len(arquivos) == 1
    doc = Document(arquivos[0])

    # Verifica texto e tabelas do documento gerado
    all_text = " ".join([p.text for p in doc.paragraphs])
    assert "QUADRO 1 – NÃO CONFORMIDADES IDENTIFICADAS CRC - 27/05/2026" in all_text

    table_texts = [" ".join([c.text for row in t.rows for c in row.cells]) for t in doc.tables]
    combined_tables = " ".join(table_texts)

    assert "CRC.SH015.0646+0648/2026.001" in combined_tables
    assert "CRC.SH015.0646+0648/2026.002" in combined_tables
    assert "Fissuras" in combined_tables or "FI" in combined_tables
    assert "Tachões soltos" in combined_tables
    print("[PASS] test_gerar_relatorio_crc_pipeline")
